from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path


OUT = Path('/Users/sushi/Downloads/新一代ERP产品与技术架构规划建议书.docx')
NAVY = '051C2C'
CYAN = '00A9F4'
INK = '222222'
GRAY = '666666'
LIGHT = 'F2F4F7'
PALE = 'EAF6FC'
WHITE = 'FFFFFF'


def rgb(hexstr):
    return RGBColor.from_string(hexstr)


def set_run(run, size=11, bold=False, color=INK, italic=False, east='Hiragino Sans GB'):
    run.font.name = 'Hiragino Sans GB'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), east)
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Hiragino Sans GB')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Hiragino Sans GB')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')


def set_table_geometry(table, widths):
    total = sum(widths)
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(total)); tblW.set(qn('w:type'), 'dxa')
    tblInd = tblPr.find(qn('w:tblInd'))
    if tblInd is None:
        tblInd = OxmlElement('w:tblInd'); tblPr.append(tblInd)
    tblInd.set(qn('w:w'), '120'); tblInd.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(width)); grid.append(gc)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tcW = cell._tc.get_or_add_tcPr().find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW'); cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn('w:w'), str(widths[idx])); tcW.set(qn('w:type'), 'dxa')
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)


def add_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'; table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]; set_cell_shading(c, NAVY)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.05
        set_run(p.add_run(h), font_size, True, WHITE)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.05
            if i == 0 and len(headers) > 2: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run(p.add_run(str(val)), font_size, False, INK)
        if len(table.rows) % 2 == 1:
            for c in cells: set_cell_shading(c, 'F8FAFC')
    set_table_geometry(table, widths); set_repeat_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_p(doc, text='', bold_prefix=None, align=None, italic=False):
    p = doc.add_paragraph(style='Normal')
    if align is not None: p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        set_run(p.add_run(bold_prefix), 11, True, NAVY)
        set_run(p.add_run(text[len(bold_prefix):]), 11, False, INK, italic)
    else:
        set_run(p.add_run(text), 11, False, INK, italic)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    set_run(p.add_run(text), 10.7, False, INK)
    return p


def number(doc, text):
    p = doc.add_paragraph(style='List Number')
    set_run(p.add_run(text), 10.7, False, INK)
    return p


def callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1); table.style = 'Table Grid'
    set_table_geometry(table, [9360]); c = table.cell(0, 0); set_cell_shading(c, PALE)
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(label + '  '), 11, True, NAVY)
    set_run(p.add_run(text), 11, False, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def page_break(doc):
    doc.add_page_break()


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

# Narrative proposal preset
styles = doc.styles
normal = styles['Normal']; normal.font.name = 'Hiragino Sans GB'; normal.font.size = Pt(11)
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Hiragino Sans GB'); normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Hiragino Sans GB'); normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Hiragino Sans GB')
normal.paragraph_format.space_before = Pt(0); normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.333; normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name, size, before, after, color in [
    ('Heading 1',16,18,10,NAVY),('Heading 2',13,12,6,NAVY),('Heading 3',12,8,4,'1F4D78')]:
    st=styles[name]; st.font.name='Hiragino Sans GB'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=rgb(color)
    st._element.rPr.rFonts.set(qn('w:ascii'), 'Hiragino Sans GB'); st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Hiragino Sans GB'); st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Hiragino Sans GB')
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
for name in ['List Bullet','List Bullet 2','List Number']:
    st=styles[name]; st.font.name='Hiragino Sans GB'; st.font.size=Pt(10.7); st._element.rPr.rFonts.set(qn('w:ascii'),'Hiragino Sans GB'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Hiragino Sans GB'); st._element.rPr.rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB')
    st.paragraph_format.space_after=Pt(4); st.paragraph_format.line_spacing=1.208

# Running header/footer
hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_run(hp.add_run('新一代 ERP 产品与技术架构规划建议书'), 8.5, False, GRAY)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run(fp.add_run('内部评审稿  |  '), 8.5, False, GRAY)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); fp._p.append(field)

# Cover
for _ in range(5): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
set_run(p.add_run('产品与技术规划建议书'), 12, True, CYAN)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(8)
set_run(p.add_run('新一代 ERP'), 30, True, NAVY)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(22)
set_run(p.add_run('产品能力、数据对象、服务化与集成架构'), 16, False, NAVY)
callout(doc, '核心主张', '以少量稳定核心对象统一语义，以中等粒度 PBC 承载业务，以事件和 API 解耦，以渐进式微服务拆分同时实现扩展性、非功能指标与成本控制。')
for _ in range(4): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
set_run(p.add_run('依据：《产品需求分析报告》与《Next-Gen ERP Proposal Insights》'), 10, False, GRAY)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
set_run(p.add_run('版本：V1.0  |  2026年8月'), 10, False, GRAY)

page_break(doc)
doc.add_heading('执行摘要', level=1)
add_p(doc, '本建议书面向新一代 ERP 的产品规划和技术架构决策。综合两份输入材料，建议不把现有 11 个子产品、52 个应用模块、L4 应用服务、业务对象、微服务和智能体建立一一对应关系，而是通过统一业务语义、能力域重组和渐进式部署，形成可组合、可治理、可演进的企业业务平台。')
callout(doc, '结论', '减少对象数量、实现微服务快速扩展、获得高 API 非功能指标并降低总体成本是可以兼容的；关键是减少重复概念和调用链，而不是把所有业务压缩成万能对象，也不是一开始建设大量微服务。')
doc.add_heading('建议管理层批准的六项决策', level=2)
for x in [
    '将约 52 个应用模块重组为 12–16 个业务能力域，首期仅形成 8–12 个独立部署单元。',
    '建立约 14–18 类企业核心对象，以稳定内核、类型化扩展和明确数据所有权消除重复对象。',
    '标品覆盖的人力、财务、采购、供应和项目能力优先复用，通过 API、防腐层和事件提升为 PBC。',
    '差异化投入集中于工器具、仓储运输、设备作业、配置变更、安全质量环境、金蝶迁移与协同。',
    '建立统一 API 网关、事件平台、持久化流程、权限、元数据、审计和可观测平台，避免重复建设。',
    'Agent、MCP 和生成式 UI 建立在稳定业务 API 之上；核算、计税、库存、支付等确定性规则不交给大模型。']:
    number(doc, x)

doc.add_heading('预期规划结果', level=2)
add_table(doc, ['规划维度','建议目标','价值'], [
    ['核心对象','约14–18类','减少跨域同义对象和映射成本'],
    ['业务能力域','12–16个','形成稳定产品边界与投资单元'],
    ['首期部署单元','8–12个','控制分布式系统复杂度'],
    ['关键接口','分级SLO','把高可用预算投入关键交易'],
    ['建设策略','复用优先','扩大已有标品与平台投资收益'],
], [1800,2100,5460])

page_break(doc)
doc.add_heading('一、项目背景与规划原则', level=1)
doc.add_heading('1.1 需求范围', level=2)
add_p(doc, '现有需求覆盖企业经营与生产运营两大产品系列，包括人力资源、财务、采购、供应、项目、设备可靠性、运行、作业、配置、安全质量环境和公共配置。公共配置还承担统一门户、基础数据、安全、运维、应用连接和数据迁移。')
add_p(doc, '结合最新洞察，目标产品还需要支持 PBC、API First、事件驱动、SOP-as-Code、多智能体、语义与情节记忆以及生成式 UI。但这些技术能力必须围绕业务价值分阶段引入，不能反向驱动业务碎片化。')
doc.add_heading('1.2 规划原则', level=2)
for label,text in [
    ('投资利用率最大化：','已有标品能覆盖的能力优先复用；公共技术能力一次建设、多域复用；开源项目优先作为平台组件而不是复制业务系统。'),
    ('高内聚低耦合：','服务边界按业务不变量、数据所有权和团队责任确定，不按菜单、数据表或单据名称拆分。'),
    ('语义统一、存储自治：','统一对象定义、编码和事件语义，但禁止多个服务直接写同一业务表。'),
    ('渐进服务化：','先形成清晰逻辑边界，再根据吞吐、可用性、发布节奏和组织自治逐步物理拆分。'),
    ('确定性优先：','财务、税务、库存、权限等规则由确定性服务执行，AI 负责意图、辅助判断和受控编排。'),
    ('分级非功能：','不要求所有接口达到同一最高等级，以业务关键性分配可用性和性能成本。')]:
    add_p(doc, label+text, bold_prefix=label)

doc.add_heading('1.3 关键边界调整', level=2)
bullet(doc, '本建议书不包含特定行业专用安全场景及其专有业务设计。')
bullet(doc, '保留通用设备、工器具、点巡检、作业许可、安全、质量和环境管理。')
bullet(doc, '存量 ERP 集成对象统一按金蝶规划，包括迁移、共存、对账和切换。')

page_break(doc)
doc.add_heading('二、现状诊断与核心问题', level=1)
doc.add_heading('2.1 当前需求结构的主要风险', level=2)
add_table(doc, ['问题','表现','潜在后果'], [
    ['层级混用','模块、应用服务、对象、接口边界接近一一对应','服务数量膨胀，边界不稳定'],
    ['对象重复','计划、订单、单据、组织等概念在多域分别定义','映射多、数据口径不一致'],
    ['分析能力分散','多个子产品各自建设“智能分析”','平台和指标重复投资'],
    ['公共能力过宽','基础数据、连接、迁移、安全、运维混在一个产品中','职责不清，形成平台单体'],
    ['同步依赖过多','端到端业务容易形成跨域同步调用链','延迟放大、故障级联'],
    ['AI 目标超前','Agent、MCP、A2A 先于稳定 API 和治理','不可控调用和重复封装'],
], [1700,3600,4060])

doc.add_heading('2.2 不应采用的拆分方式', level=2)
bullet(doc, '一个 L3 模块对应一个微服务。模块是产品导航和责任视图，不天然等于运行边界。')
bullet(doc, '一个业务对象对应一个微服务。对象之间存在强业务不变量时应由同一能力域维护。')
bullet(doc, '一个 Agent 对应一个微服务。Agent 是调用者和编排者，不应拥有底层交易真相。')
bullet(doc, '所有服务共用一个数据库。共享库会使发布、扩容和故障隔离名存实亡。')
bullet(doc, '把所有业务统一成动态 JSON 对象。对象数量表面减少，约束、索引、权限和契约成本却显著上升。')

callout(doc, '架构判断', '真正需要控制的是核心概念数量、部署单元数量和跨服务调用数量，而不是单纯追求微服务数量或数据库表数量。')

page_break(doc)
doc.add_heading('三、目标产品能力蓝图', level=1)
add_p(doc, '建议将产品能力划分为企业核心业务、公共平台、数据智能和智能协作四层。产品能力域是投资、路线图和责任边界；部署单元是运行边界，两者允许一对多或多对一。')
add_table(doc, ['能力域','整合范围','规划策略'], [
    ['企业与组织','组织、人员、岗位、客户、供应商、承包商','统一 Party 与 Assignment 模型'],
    ['人力运营','入转调离、合同、考勤、薪酬、派遣','标品优先，规则独立治理'],
    ['商业与合同','询报价、框架协议、采购与客户协议','统一 Agreement 语义'],
    ['需求与订单','采购需求、销售订单、采购订单、工作任务','共享内核、领域所有'],
    ['计划与资源','需求、采购、制造、项目、维修计划','统一 Plan 结构，算法可插拔'],
    ['履行与物流','验收、收发货、仓储、运输、库存','事件驱动，按吞吐扩展'],
    ['制造与作业','制造订单、工单、任务清单、操作单','统一 Work 模型族'],
    ['资产与设备','资产、设备、功能位置、维修、工器具','统一 Resource/Asset 内核'],
    ['项目管理','立项、WBS、进度、变更、验收','复用计划、采购与财务'],
    ['财务核算','应收、应付、成本、资产、税务、总账','强一致内核，谨慎拆分'],
    ['预算与控制','企业预算、项目预算、额度与信用控制','建设统一控制能力'],
    ['安全质量环境','作业许可、检查、环境监测、整改','统一策略、证据与审计'],
    ['数据与分析','指标、报告、预测、质量、血缘','与交易解耦，平台化建设'],
    ['智能协作','Agent、SOP、记忆、生成式 UI','构建于受控 API 之上'],
], [1800,4200,3360], 8.7)

doc.add_heading('3.1 建设、复用与采购边界', level=2)
add_table(doc, ['策略','能力范围','建议'], [
    ['复用标品','组织、人事、考勤、薪酬、应收应付、核算、预算、采购、供应计划、项目','通过稳定 API、事件和防腐层提升为 PBC，不复制底层规则'],
    ['重点自研','工器具、仓储运输、设备作业、配置变更、安全质量环境、差异化分析','共享公共对象、流程、权限、规则和事件能力'],
    ['平台共建','网关、事件、工作流、权限、元数据、审计、观测、AI治理','统一建设一次，禁止各子产品重复选型'],
], [1500,3900,3960])

page_break(doc)
doc.add_heading('四、企业核心数据对象规划', level=1)
doc.add_heading('4.1 建议的核心对象', level=2)
add_table(doc, ['对象','覆盖范围','主要责任域'], [
    ['Party','企业、组织、人员、客户、供应商、承包商','企业与组织'],
    ['Role & Assignment','岗位、职位、职级、雇佣、组织与汇报关系','企业与组织/人力'],
    ['Resource','物料、产品、设备、资产、工器具、服务项','主数据/资产'],
    ['Resource Structure','BOM、设备结构、功能位置、工艺路线、WBS','供应/资产/项目'],
    ['Location','工厂、仓库、库位、作业地点、运输节点','履行与物流'],
    ['Agreement','劳动合同、框架协议、采购合同、客户与项目协议','商业与合同'],
    ['Demand','采购需求、物料需求、服务需求、项目需求','需求与订单'],
    ['Plan','预算、采购、制造、维修、项目和运输计划','计划与资源'],
    ['Order / Work','采购、销售、制造订单，工单与任务','订单/制造/作业'],
    ['Fulfillment','收货、验收、发运、调拨、退货、完工','履行与物流'],
    ['Inventory Position','余额、可用量、预留、在途量','库存'],
    ['Financial Transaction','发票、收付款、成本、资产与税务交易','财务核算'],
    ['Accounting Entry','分录、凭证、余额与期间','财务核算'],
    ['Control','预算、信用、额度、审批策略','预算与控制'],
    ['Case / Change','项目变更、配置变更、异常、纠正行动','项目/配置/质量'],
    ['Permit & Safety Control','作业许可、安全措施、风险与环境控制','安全质量环境'],
    ['Observation','点巡检、质量安全检查、设备测量、环境监测','设备/质量环境'],
    ['Business Event','创建、审批、下达、发运、入库、过账等事实','企业事件平台'],
], [1900,4600,2860], 8.5)

doc.add_heading('4.2 对象减少的实现机制', level=2)
for x in [
    '对象族统一：采购订单、销售订单、制造订单和工单共享 Order/Work 的身份、参与方、状态、行项目和履行关系，但保留领域扩展。',
    '类型与状态收敛：标准发票、借项发票、贷项发票作为 Invoice 的不同类型，不建立三套顶层对象。',
    '主体与业务状态分离：Resource 统一标识物料、设备、资产和工具；库存、维修、折旧由相应能力域拥有。',
    '引用替代复制：跨域只保存全局 ID、必要快照和契约字段，不复制完整主对象。',
    '事件表达变化：状态变化发布为不可变业务事件，分析和集成不直接读取交易库。']:
    number(doc, x)

callout(doc, '边界原则', '共享对象定义不等于共享数据库。每个对象必须有唯一记录源、唯一写入责任域和明确的数据契约。')

page_break(doc)
doc.add_heading('五、PBC、微服务与部署单元设计', level=1)
doc.add_heading('5.1 推荐首期部署拓扑', level=2)
add_table(doc, ['部署单元','包含的能力','后续拆分触发条件'], [
    ['组织与人力服务','企业组织、人事、合同、考勤、薪酬适配','薪酬周期负载或独立合规要求'],
    ['商业采购服务','供应商、询报价、协议、采购需求与订单','寻源与履行发布节奏明显分化'],
    ['计划与供应服务','计划、销售订单、制造、库存协调','计划计算和交易吞吐出现差异'],
    ['仓储与物流服务','仓储作业、收发货、运输','高并发或现场独立部署'],
    ['资产与作业服务','设备、维修、工器具、工单、巡检','设备数据量或边缘场景增长'],
    ['项目服务','项目结构、计划、进度、变更与验收','大型项目形成独立产品团队'],
    ['财务与控制服务','应收应付、核算、预算、成本、税务适配','结算、报表或预算出现独立扩展压力'],
    ['安全质量环境服务','许可、检查、隐患、整改、环境控制','法规或现场隔离需求'],
    ['数据智能服务','指标、报告、预测、语义、质量','分析规模与交易平台独立扩容'],
    ['集成与智能协作','金蝶适配、API、事件、SOP、Agent 工具','Agent流量或集成流量独立增长'],
], [2200,4300,2860], 8.6)

doc.add_heading('5.2 服务拆分判据', level=2)
add_p(doc, '只有满足下列一个或多个条件时，才将逻辑 PBC 拆为独立微服务：')
for x in ['需要独立扩容或资源类型明显不同；','可用性与故障隔离等级明显不同；','发布频率和团队责任长期独立；','存在清晰业务不变量和数据所有权；','安全、法规或地域部署边界不同。']:
    bullet(doc, x)
add_p(doc, '如果拆分后绝大多数请求仍必须同步调用另一个服务、仍共享数据库、必须联动发布，说明边界并未成熟。')

doc.add_heading('5.3 事务和一致性', level=2)
bullet(doc, '服务内部使用本地 ACID 事务维护业务不变量。')
bullet(doc, '跨服务使用 Outbox、幂等消费者、Saga/补偿和持久化流程。')
bullet(doc, '记账、付款、库存扣减等关键命令必须携带幂等键和业务版本。')
bullet(doc, '分析查询通过事件生成读模型，避免在线跨服务联表。')

page_break(doc)
doc.add_heading('六、目标集成架构与金蝶迁移', level=1)
doc.add_heading('6.1 总体架构', level=2)
add_p(doc, '接入层通过统一 API/AI Gateway 提供认证、限流、路由和审计；场景 API/BFF 负责面向门户、移动端和 Agent 的交互聚合；PBC 维护业务逻辑和私有数据；企业事件平台传播业务事实；流程引擎承载跨域长事务；元数据和权限平台提供横向治理。')
add_table(doc, ['层次','核心职责','关键约束'], [
    ['接入与体验','门户、移动端、生成式UI、外部应用','不直接访问领域数据库'],
    ['API/AI网关','认证、限流、路由、灰度、审计、协议转换','内外接口分层，策略集中'],
    ['场景API/BFF','查询聚合、交互适配、命令受理','不复制核心业务规则'],
    ['PBC能力层','业务规则、数据所有权、领域事件','高内聚、私有写模型'],
    ['流程与事件','长事务、异步协作、重试、补偿','事件不可变、消费者幂等'],
    ['数据与智能','读模型、指标、语义、记忆、AI评估','与在线交易解耦'],
    ['治理与观测','权限、契约、血缘、追踪、审计','统一平台、全链路可追溯'],
], [1700,4500,3160])

doc.add_heading('6.2 金蝶集成策略', level=2)
add_p(doc, '金蝶在迁移期作为存量 ERP 和部分业务数据的记录源。新一代 ERP 不直接复制金蝶表结构，而是在集成边界建立金蝶防腐适配层，将组织、供应商、客户、物料、科目、单据和状态转换为统一业务对象与事件。')
for x in [
    '建立金蝶组织、人员、客户、供应商、物料、会计科目等主数据映射。',
    '适配采购、库存、应收、应付、总账等业务接口，并建立调用限流、重试和熔断。',
    '建立金蝶单据号与新 ERP 全局对象 ID 对照，保留来源系统与迁移批次。',
    '建设外围作业系统—金蝶—新一代 ERP 的集成三通，支持并行验证。',
    '通过增量同步、差异检测、日终对账和异常工单控制共存期一致性。',
    '按业务域逐步切换系统记录源；切换后关闭金蝶对应写接口，避免长期双写。']:
    number(doc, x)

doc.add_heading('6.3 迁移波次', level=2)
add_table(doc, ['波次','范围','退出条件'], [
    ['波次0','对象标准、编码、数据质量、接口盘点','核心对象和责任域获批'],
    ['波次1','主数据与只读查询','完整率、唯一性、映射准确率达标'],
    ['波次2','低风险业务域和异步事件','双系统对账稳定'],
    ['波次3','采购、库存、项目等核心交易','关键链路压测和回退演练通过'],
    ['波次4','财务结算和最终切换','关账、余额、凭证和审计核验通过'],
], [1300,4600,3460])

page_break(doc)
doc.add_heading('七、API 非功能指标与成本控制', level=1)
doc.add_heading('7.1 分级 SLO', level=2)
add_table(doc, ['等级','典型接口','建议目标','实现重点'], [
    ['S0 核心交易','记账、支付、库存扣减、权限校验','可用性≥99.99%；严格幂等','隔离部署、强审计、容量冗余'],
    ['S1 在线业务','订单、审批、主数据查询','可用性≥99.95%；P95≤200ms','缓存、读模型、限流、降级'],
    ['S2 复杂命令','履行、关账、计划运算','300–500ms内返回受理状态','异步流程、进度查询、补偿'],
    ['S3 分析智能','报表、预测、Agent推理','异步或流式；独立配额','结果缓存、批处理、模型路由'],
], [1300,2800,2800,2460], 8.8)
add_p(doc, '以上为建议基线，最终值必须通过真实业务量、峰值并发、部署环境和容灾等级校准。API 网关厂商或开源项目提供的基准测试不能替代本项目压测。')

doc.add_heading('7.2 同时实现高性能与低成本的措施', level=2)
for x in [
    '逻辑服务化、物理渐进拆分：首期一个部署单元可承载多个高内聚 PBC。',
    '同步查询、异步履约：跨域长流程快速返回受理状态，避免线程和连接长期占用。',
    '减少调用放大：场景 API 聚合读取；批量接口替代循环调用；热点数据使用读模型和缓存。',
    '平台能力共用：网关、消息、权限、流程、观测、审计、密钥和配置只建设一套。',
    '弹性按压力点拆分：计划计算、报表、Agent 等资源密集型能力独立扩展，不带动核心交易整体扩容。',
    'Agent 成功路径固化：反复验证的高频路径转为确定性流程，减少模型 Token 和决策延迟。']:
    bullet(doc, x)

doc.add_heading('7.3 建议纳入验收的工程指标', level=2)
add_table(doc, ['类别','指标'], [
    ['可靠性','错误预算、恢复时间、消息积压、补偿成功率、灾备切换时间'],
    ['性能','P50/P95/P99延迟、吞吐、并发、数据库与缓存命中率'],
    ['质量','契约兼容率、幂等覆盖率、事件重复与丢失检测、数据对账差异'],
    ['安全','细粒度授权命中、敏感数据脱敏、审计完整率、Agent越权拦截'],
    ['成本','单笔交易成本、单流程资源消耗、平台固定成本、模型调用成本'],
], [1800,7560])

page_break(doc)
doc.add_heading('八、开源项目复用建议', level=1)
add_p(doc, '开源选型应以能力缺口、成熟度、活跃度、许可、运维复杂度和替换成本综合判断。GitHub 热度只作为生态活跃度信号，不作为采购或生产准入结论。')
add_table(doc, ['项目','建议用途','引入建议与风险'], [
    ['Dapr','服务调用、Pub/Sub、状态、工作流、密钥与观测抽象','建议PoC；可减少通用分布式代码，但需评估Sidecar资源和团队运维能力'],
    ['Temporal','长事务、SOP、重试、补偿和持久化执行','优先验证采购履行、入转调离、项目验收和关账场景'],
    ['NATS / 现有Kafka','企业事件总线','二选一；已有Kafka能力时优先复用，避免新增消息栈'],
    ['Apache APISIX','API/AI网关、鉴权、限流、灰度、观测、MCP桥接','建议PoC；必须以真实流量和安全基线验证'],
    ['OpenFGA','关系型细粒度权限','适合组织、项目、数据域和Agent工具权限；需与现有IAM集成'],
    ['OpenMetadata / DataHub','元数据、术语、血缘、质量和契约','二选一；避免同时全面建设两个平台'],
    ['MCP SDK/参考服务器','Agent工具接口标准与开发参考','参考服务器不是生产就绪成品，需补充鉴权、审计、限流和输入校验'],
    ['A2A','未来Agent间互操作','暂缓到跨域Agent网络形成后；不作为一期前置条件'],
    ['ERPNext / Apache OFBiz','业务模型与原型参考','不建议替换现有标品；ERPNext为GPL-3.0，商业嵌入需法律评估'],
], [1700,3100,4560], 8.5)

doc.add_heading('8.1 推荐的最小组合', level=2)
bullet(doc, 'API Gateway：APISIX 或现有企业网关二选一。')
bullet(doc, '微服务通用能力：Dapr PoC 后决定是否采用。')
bullet(doc, '流程：Temporal，优先覆盖跨域长流程。')
bullet(doc, '消息：复用现有 Kafka；没有既有能力时评估 NATS。')
bullet(doc, '权限：现有 IAM＋OpenFGA 细粒度授权。')
bullet(doc, '元数据：OpenMetadata 与 DataHub 二选一。')
bullet(doc, '观测：OpenTelemetry 统一标准。')

page_break(doc)
doc.add_heading('九、Agentic ERP 建设边界', level=1)
doc.add_heading('9.1 Agent 的合理职责', level=2)
bullet(doc, '自然语言意图识别、任务拆解和场景路由。')
bullet(doc, '调用经过批准、版本受控且范围最小化的业务工具。')
bullet(doc, '辅助分析异常、生成解释、提出建议和组织审批材料。')
bullet(doc, '在 SOP 图和策略约束下编排跨域任务。')
bullet(doc, '记录执行轨迹、业务结果和人工反馈，用于评估改进。')

doc.add_heading('9.2 不应交给 Agent 的职责', level=2)
bullet(doc, '直接修改核心数据库或绕过业务 API。')
bullet(doc, '自行决定会计分录、税务处理、库存扣减或支付规则。')
bullet(doc, '在没有授权模型、审计和人工复核的情况下执行高风险动作。')
bullet(doc, '通过 Prompt 代替确定性校验、业务规则和数据完整性约束。')

doc.add_heading('9.3 MCP、A2A 与 SOP-as-Code', level=2)
add_table(doc, ['机制','定位','采用条件'], [
    ['MCP','Agent到工具/资源的标准适配层','业务API稳定，鉴权、审计、输入输出Schema完备'],
    ['A2A','Agent之间的能力发现与协作协议','出现跨团队、跨平台Agent网络和互操作需求'],
    ['SOP-as-Code','将流程、条件、权限和人工节点显式化','流程规则可验证、版本化、回放和审计'],
    ['双态记忆','语义规则与历史情节分层存储','完成脱敏、保留周期、质量评分和访问控制'],
], [1600,3900,3860])
callout(doc, '治理要求', '第二份输入材料中的准确率和ROI数字应作为待验证假设。试点必须建立人工基线、样本集、失败分类和业务结果指标，避免用模型演示替代生产证明。')

page_break(doc)
doc.add_heading('十、实施路线图', level=1)
add_table(doc, ['阶段','周期','重点工作','关键产出'], [
    ['阶段0：规划与基线','0–2个月','对象盘点、能力地图、系统记录源、SLO与成本基线','核心对象V1、PBC目录、决策记录、PoC范围'],
    ['阶段1：平台与试点','3–6个月','网关、事件、流程、权限、观测；金蝶防腐层；三个端到端场景','平台最小闭环、契约体系、压测与安全报告'],
    ['阶段2：领域迁移','7–12个月','采购、供应、项目、设备作业等按波次迁移','8–12个部署单元、双系统对账、回退机制'],
    ['阶段3：财务与智能','13–18个月','财务切换、SOP、受控Agent、读模型和指标平台','核心交易闭环、Agent生产试点、Outcome评估'],
    ['阶段4：优化扩展','19–24个月','按压力拆分服务、性能成本优化、生态集成','弹性扩展、服务目录、持续改进机制'],
], [1600,1200,3800,2760], 8.5)

doc.add_heading('10.1 首批三个验证场景', level=2)
for label,text in [
    ('采购需求到付款：','验证需求、订单、验收、发票、预算控制、事件和跨域流程。'),
    ('项目到资产：','验证项目结构、计划、采购、进度、成本、验收和资产转固。'),
    ('设备异常到工单闭环：','验证设备、观察、工单、工器具、作业许可、整改和分析。')]:
    add_p(doc, label+text, bold_prefix=label)

doc.add_heading('10.2 阶段门', level=2)
bullet(doc, '对象门：核心对象、唯一记录源、ID 和事件语义获业务与技术共同批准。')
bullet(doc, '架构门：关键场景完成故障、重试、幂等、补偿和回退验证。')
bullet(doc, '性能门：真实数据量下达到分级 SLO，且容量成本可解释。')
bullet(doc, '迁移门：金蝶与新 ERP 对账差异在阈值内，业务连续性演练通过。')
bullet(doc, 'AI门：高风险动作全量受控，准确率、越权率和人工接管率达到试点目标。')

page_break(doc)
doc.add_heading('十一、治理机制、风险与应对', level=1)
doc.add_heading('11.1 建议的治理组织', level=2)
add_table(doc, ['治理主体','职责'], [
    ['产品能力委员会','批准能力边界、路线图、复用与自研决策'],
    ['企业对象委员会','维护核心对象、编码、术语、责任域和数据契约'],
    ['架构评审委员会','控制服务拆分、技术栈数量、NFR和例外'],
    ['API与事件委员会','治理版本、兼容性、幂等、事件语义和弃用'],
    ['AI治理委员会','批准Agent工具、SOP、模型、风险等级和评估基线'],
], [2600,6760])

doc.add_heading('11.2 主要风险', level=2)
add_table(doc, ['风险','影响','应对'], [
    ['过度对象收敛','形成万能对象，约束和性能恶化','稳定内核＋类型扩展；限制动态字段使用范围'],
    ['过早微服务化','成本、故障点和调用链激增','首期8–12个部署单元；以拆分判据控制'],
    ['金蝶长期双写','数据冲突和业务责任不清','按域确定唯一记录源；切换后关闭旧写入口'],
    ['开源栈过多','运维和升级负担增加','每类能力原则上只选一个主平台'],
    ['同步集成过重','级联故障与延迟放大','异步事件、受理模式、读模型和降级'],
    ['Agent越权或幻觉','错误交易、合规和审计风险','最小权限、规则校验、HITL、全轨迹审计'],
    ['指标承诺失真','立项收益无法兑现','所有收益数字建立基线、样本和归因方法'],
], [2100,3000,4260], 8.7)

page_break(doc)
doc.add_heading('十二、最终建议与立项决策', level=1)
callout(doc, '最终判断', '可以在减少核心对象数量的同时，实现微服务快速扩展、API高非功能指标和较低成本；实现路径是“语义收敛、能力解耦、部署渐进、接口分级、平台复用”。')
add_p(doc, '建议批准以“核心对象与能力边界优先”为原则推进新一代 ERP，而不是以微服务数量或 AI 功能数量衡量现代化程度。对象模型保持少而稳定，产品能力保持可组合，部署单元根据真实压力拆分，关键 API 获得更高等级保障，普通接口按业务价值控制投入。')
doc.add_heading('建议立即启动的工作', level=2)
for x in [
    '在四周内完成现有对象、模块、接口和金蝶数据源盘点。',
    '成立产品能力、企业对象、架构和AI治理四类联合工作组。',
    '确定三个端到端试点场景及现状性能、成本和人工工作量基线。',
    '完成 API 网关、流程、事件、权限和元数据平台的最小 PoC 选型。',
    '形成核心对象V1、PBC目录V1、系统记录源矩阵和金蝶迁移波次计划。',
    '在试点验证通过前，不启动大规模微服务拆分和跨域Agent网络建设。']:
    number(doc, x)

doc.add_heading('附录A：来源与参考项目', level=1)
refs = [
    '《新一代ERP产品设计技术支持服务项目_产品需求分析报告》，2024年10月。',
    '《Next-Gen ERP Proposal Insights》。',
    'Dapr：https://github.com/dapr/dapr',
    'Temporal：https://github.com/temporalio/temporal',
    'NATS：https://github.com/nats-io/nats-server',
    'Apache APISIX：https://github.com/apache/apisix',
    'OpenFGA：https://github.com/openfga/openfga',
    'OpenMetadata：https://github.com/open-metadata/OpenMetadata',
    'DataHub：https://github.com/datahub-project/datahub',
    'Model Context Protocol Servers：https://github.com/modelcontextprotocol/servers',
    'A2A Protocol：https://github.com/a2aproject/A2A',
    'ERPNext：https://github.com/frappe/erpnext',
]
for r in refs: bullet(doc, r)

# Document properties
doc.core_properties.title = '新一代ERP产品与技术架构规划建议书'
doc.core_properties.subject = '产品能力、数据对象、服务化与金蝶集成架构'
doc.core_properties.author = 'Codex'
doc.save(OUT)
print(OUT)
